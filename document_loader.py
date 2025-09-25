#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Загрузчик документов в базу данных Rubin AI v2.0
"""

import os
import sys
import sqlite3
import hashlib
import mimetypes
from pathlib import Path
import logging
from datetime import datetime
import json

class DocumentLoader:
    def __init__(self, db_path="rubin_ai_documents.db"):
        self.db_path = db_path
        self.setup_logging()
        self.setup_database()
        
    def setup_logging(self):
        """Настройка логирования"""
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler('document_loader.log', encoding='utf-8'),
                logging.StreamHandler(sys.stdout)
            ]
        )
        self.logger = logging.getLogger(__name__)
        
    def setup_database(self):
        """Создание таблиц базы данных"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Таблица документов
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_path TEXT UNIQUE NOT NULL,
                    file_name TEXT NOT NULL,
                    file_size INTEGER,
                    file_type TEXT,
                    file_hash TEXT,
                    content TEXT,
                    metadata TEXT,
                    category TEXT,
                    tags TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            ''')
            
            # Таблица индекса для поиска
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS document_index (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER,
                    keyword TEXT,
                    position INTEGER,
                    context TEXT,
                    FOREIGN KEY (document_id) REFERENCES documents (id)
                )
            ''')
            
            # Таблица категорий
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS categories (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT UNIQUE NOT NULL,
                    description TEXT,
                    parent_id INTEGER,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            ''')
            
            conn.commit()
            conn.close()
            
            self.logger.info("✅ База данных инициализирована")
            
        except Exception as e:
            self.logger.error(f"❌ Ошибка инициализации базы данных: {e}")
            
    def get_file_hash(self, file_path):
        """Получение хеша файла"""
        try:
            hash_md5 = hashlib.md5()
            with open(file_path, "rb") as f:
                for chunk in iter(lambda: f.read(4096), b""):
                    hash_md5.update(chunk)
            return hash_md5.hexdigest()
        except Exception as e:
            self.logger.error(f"Ошибка получения хеша файла {file_path}: {e}")
            return None
            
    def extract_text_content(self, file_path):
        """Извлечение текстового содержимого из файла"""
        try:
            file_type = mimetypes.guess_type(file_path)[0]
            
            if file_type is None:
                # Попробуем определить по расширению
                ext = Path(file_path).suffix.lower()
                if ext in ['.txt', '.md', '.rst']:
                    file_type = 'text/plain'
                elif ext in ['.pdf']:
                    file_type = 'application/pdf'
                elif ext in ['.doc', '.docx']:
                    file_type = 'application/msword'
                    
            if file_type and file_type.startswith('text/'):
                # Текстовые файлы
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
                    
            elif file_type == 'application/pdf':
                # PDF файлы
                try:
                    import PyPDF2
                    with open(file_path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        text = ""
                        for page in reader.pages:
                            text += page.extract_text() + "\n"
                        return text
                except ImportError:
                    self.logger.warning("PyPDF2 не установлен, пропускаем PDF файл")
                    return None
                    
            elif file_type in ['application/msword', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document']:
                # Word документы
                try:
                    import docx
                    doc = docx.Document(file_path)
                    text = ""
                    for paragraph in doc.paragraphs:
                        text += paragraph.text + "\n"
                    return text
                except ImportError:
                    self.logger.warning("python-docx не установлен, пропускаем Word документ")
                    return None
                    
            else:
                self.logger.warning(f"Неподдерживаемый тип файла: {file_type}")
                return None
                
        except Exception as e:
            self.logger.error(f"Ошибка извлечения текста из {file_path}: {e}")
            return None
            
    def categorize_document(self, file_path, content):
        """Автоматическая категоризация документа"""
        file_name = Path(file_path).name.lower()
        content_lower = content.lower() if content else ""
        
        # Ключевые слова для категоризации
        categories = {
            'Электротехника': ['электричество', 'ток', 'напряжение', 'схема', 'резистор', 'конденсатор', 'индуктивность', 'мощность', 'закон ома', 'кирхгоф'],
            'Программирование': ['программирование', 'код', 'алгоритм', 'функция', 'переменная', 'цикл', 'условие', 'python', 'java', 'c++', 'javascript'],
            'Автоматизация': ['автоматизация', 'plc', 'scada', 'контроллер', 'датчик', 'исполнительный', 'регулирование', 'промышленность'],
            'Радиотехника': ['радио', 'антенна', 'частота', 'модуляция', 'передатчик', 'приемник', 'сигнал', 'волна'],
            'Механика': ['механика', 'движение', 'сила', 'масса', 'скорость', 'ускорение', 'кинематика', 'динамика'],
            'Математика': ['математика', 'уравнение', 'функция', 'производная', 'интеграл', 'геометрия', 'алгебра', 'тригонометрия']
        }
        
        # Подсчет совпадений
        category_scores = {}
        for category, keywords in categories.items():
            score = 0
            for keyword in keywords:
                if keyword in file_name:
                    score += 3  # Больший вес для имени файла
                if keyword in content_lower:
                    score += 1
            category_scores[category] = score
            
        # Возвращаем категорию с наибольшим счетом
        if category_scores:
            best_category = max(category_scores, key=category_scores.get)
            if category_scores[best_category] > 0:
                return best_category
                
        return 'Общая техническая литература'
        
    def extract_keywords(self, content):
        """Извлечение ключевых слов из содержимого"""
        if not content:
            return []
            
        # Простое извлечение ключевых слов
        words = content.lower().split()
        
        # Фильтрация стоп-слов
        stop_words = {'и', 'в', 'на', 'с', 'по', 'для', 'от', 'до', 'из', 'к', 'у', 'о', 'об', 'что', 'как', 'где', 'когда', 'почему', 'это', 'то', 'та', 'те', 'такой', 'такая', 'такое', 'такие', 'или', 'но', 'а', 'же', 'ли', 'бы', 'не', 'ни', 'уже', 'еще', 'только', 'даже', 'все', 'всё', 'всего', 'всей', 'всем', 'всеми', 'всю', 'всех', 'всего', 'всей', 'всем', 'всеми', 'всю', 'всех'}
        
        # Фильтрация коротких слов
        keywords = [word for word in words if len(word) > 3 and word not in stop_words]
        
        # Подсчет частоты
        from collections import Counter
        word_counts = Counter(keywords)
        
        # Возвращаем топ-20 ключевых слов
        return [word for word, count in word_counts.most_common(20)]
        
    def load_document(self, file_path):
        """Загрузка одного документа"""
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                self.logger.error(f"Файл не найден: {file_path}")
                return False
                
            # Получение информации о файле
            file_size = file_path.stat().st_size
            file_type = mimetypes.guess_type(str(file_path))[0] or 'unknown'
            file_hash = self.get_file_hash(file_path)
            
            # Извлечение содержимого
            content = self.extract_text_content(file_path)
            if content is None:
                self.logger.warning(f"Не удалось извлечь содержимое из {file_path}")
                return False
                
            # Категоризация
            category = self.categorize_document(file_path, content)
            
            # Извлечение ключевых слов
            keywords = self.extract_keywords(content)
            
            # Подготовка метаданных
            metadata = {
                'file_extension': file_path.suffix,
                'file_created': datetime.fromtimestamp(file_path.stat().st_ctime).isoformat(),
                'file_modified': datetime.fromtimestamp(file_path.stat().st_mtime).isoformat(),
                'keywords': keywords[:10],  # Топ-10 ключевых слов
                'word_count': len(content.split()),
                'character_count': len(content)
            }
            
            # Сохранение в базу данных
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Проверка, существует ли уже документ
            cursor.execute("SELECT id FROM documents WHERE file_path = ?", (str(file_path),))
            existing = cursor.fetchone()
            
            if existing:
                # Обновление существующего документа
                cursor.execute('''
                    UPDATE documents SET
                        file_name = ?, file_size = ?, file_type = ?, file_hash = ?,
                        content = ?, metadata = ?, category = ?, tags = ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE file_path = ?
                ''', (
                    file_path.name, file_size, file_type, file_hash,
                    content, json.dumps(metadata), category, json.dumps(keywords),
                    str(file_path)
                ))
                doc_id = existing[0]
                self.logger.info(f"🔄 Обновлен документ: {file_path.name}")
            else:
                # Вставка нового документа
                cursor.execute('''
                    INSERT INTO documents 
                    (file_path, file_name, file_size, file_type, file_hash, 
                     content, metadata, category, tags)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    str(file_path), file_path.name, file_size, file_type, file_hash,
                    content, json.dumps(metadata), category, json.dumps(keywords)
                ))
                doc_id = cursor.lastrowid
                self.logger.info(f"✅ Загружен документ: {file_path.name}")
                
            # Создание индекса для поиска
            self.create_search_index(cursor, doc_id, content)
            
            conn.commit()
            conn.close()
            
            return True
            
        except Exception as e:
            self.logger.error(f"Ошибка загрузки документа {file_path}: {e}")
            return False
            
    def create_search_index(self, cursor, doc_id, content):
        """Создание поискового индекса"""
        try:
            # Удаляем старый индекс
            cursor.execute("DELETE FROM document_index WHERE document_id = ?", (doc_id,))
            
            # Создаем новый индекс
            words = content.lower().split()
            for i, word in enumerate(words):
                if len(word) > 2:  # Индексируем только слова длиннее 2 символов
                    context = ' '.join(words[max(0, i-5):i+6])  # Контекст ±5 слов
                    cursor.execute('''
                        INSERT INTO document_index (document_id, keyword, position, context)
                        VALUES (?, ?, ?, ?)
                    ''', (doc_id, word, i, context))
                    
        except Exception as e:
            self.logger.error(f"Ошибка создания индекса: {e}")
            
    def load_directory(self, directory_path):
        """Загрузка всех документов из директории"""
        directory_path = Path(directory_path)
        
        if not directory_path.exists():
            self.logger.error(f"Директория не найдена: {directory_path}")
            return False
            
        self.logger.info(f"📁 Начинаем загрузку из директории: {directory_path}")
        
        # Поддерживаемые расширения файлов
        supported_extensions = {
            '.txt', '.md', '.rst', '.pdf', '.doc', '.docx',
            '.rtf', '.odt', '.html', '.htm', '.xml', '.json'
        }
        
        total_files = 0
        loaded_files = 0
        skipped_files = 0
        
        # Рекурсивный обход директории
        for file_path in directory_path.rglob('*'):
            if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
                total_files += 1
                
                try:
                    if self.load_document(file_path):
                        loaded_files += 1
                    else:
                        skipped_files += 1
                        
                except Exception as e:
                    self.logger.error(f"Ошибка обработки файла {file_path}: {e}")
                    skipped_files += 1
                    
        self.logger.info(f"📊 Загрузка завершена:")
        self.logger.info(f"   Всего файлов: {total_files}")
        self.logger.info(f"   Загружено: {loaded_files}")
        self.logger.info(f"   Пропущено: {skipped_files}")
        
        return loaded_files > 0
        
    def search_documents(self, query, limit=10):
        """Поиск документов по запросу"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Поиск по ключевым словам
            query_words = query.lower().split()
            
            # Поиск в индексе
            placeholders = ','.join(['?' for _ in query_words])
            cursor.execute(f'''
                SELECT DISTINCT d.id, d.file_name, d.category, d.metadata, 
                       GROUP_CONCAT(DISTINCT di.keyword) as matched_keywords
                FROM documents d
                JOIN document_index di ON d.id = di.document_id
                WHERE di.keyword IN ({placeholders})
                GROUP BY d.id
                ORDER BY COUNT(DISTINCT di.keyword) DESC
                LIMIT ?
            ''', query_words + [limit])
            
            results = cursor.fetchall()
            conn.close()
            
            return results
            
        except Exception as e:
            self.logger.error(f"Ошибка поиска: {e}")
            return []
            
    def index_document_for_vector_search(self, document_id, content, metadata=None):
        """Индексация документа для векторного поиска"""
        try:
            # Импорт здесь, чтобы избежать циклических зависимостей
            from vector_search import VectorSearchEngine
            
            vector_engine = VectorSearchEngine(self.db_path)
            success = vector_engine.index_document(document_id, content, metadata)
            
            if success:
                self.logger.info(f"✅ Документ {document_id} проиндексирован для векторного поиска")
            else:
                self.logger.warning(f"⚠️ Не удалось проиндексировать документ {document_id} для векторного поиска")
                
            return success
            
        except Exception as e:
            self.logger.error(f"❌ Ошибка векторной индексации документа {document_id}: {e}")
            return False
            
    def get_document_stats(self):
        """Получение статистики документов"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            # Общая статистика
            cursor.execute("SELECT COUNT(*) FROM documents")
            total_docs = cursor.fetchone()[0]
            
            # Статистика по категориям
            cursor.execute('''
                SELECT category, COUNT(*) as count
                FROM documents
                GROUP BY category
                ORDER BY count DESC
            ''')
            categories = cursor.fetchall()
            
            # Статистика по типам файлов
            cursor.execute('''
                SELECT file_type, COUNT(*) as count
                FROM documents
                GROUP BY file_type
                ORDER BY count DESC
            ''')
            file_types = cursor.fetchall()
            
            conn.close()
            
            return {
                'total_documents': total_docs,
                'categories': categories,
                'file_types': file_types
            }
            
        except Exception as e:
            self.logger.error(f"Ошибка получения статистики: {e}")
            return None

    def get_all_documents(self):
        """Получение всех документов из базы данных"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            
            cursor.execute('''
                SELECT id, file_name, content, 
                       COALESCE(category, '') as category, 
                       COALESCE(tags, '') as tags, 
                       COALESCE(difficulty_level, 'medium') as difficulty_level, 
                       COALESCE(file_type, '') as file_type, 
                       COALESCE(file_size, 0) as file_size, 
                       COALESCE(file_hash, '') as file_hash, 
                       created_at, updated_at
                FROM documents
                ORDER BY created_at DESC
            ''')
            
            documents = []
            for row in cursor.fetchall():
                documents.append({
                    'id': row[0],
                    'file_name': row[1],
                    'content': row[2],
                    'category': row[3],
                    'tags': row[4],
                    'difficulty_level': row[5],
                    'file_type': row[6],
                    'file_size': row[7],
                    'file_hash': row[8],
                    'created_at': row[9],
                    'updated_at': row[10],
                    'metadata': f"Category: {row[3]}, Tags: {row[4]}, Level: {row[5]}"
                })
            
            conn.close()
            return documents
            
        except Exception as e:
            self.logger.error(f"Ошибка получения всех документов: {e}")
            return []

def main():
    """Главная функция"""
    print("📚 ЗАГРУЗЧИК ДОКУМЕНТОВ RUBIN AI v2.0")
    print("=" * 50)
    
    # Путь к директории с технической литературой
    tech_literature_path = r"E:\03.Тех.литература"
    
    if not os.path.exists(tech_literature_path):
        print(f"❌ Директория не найдена: {tech_literature_path}")
        print("Пожалуйста, проверьте путь к директории с технической литературой")
        return
        
    # Создание загрузчика
    loader = DocumentLoader()
    
    # Загрузка документов
    print(f"📁 Загружаем документы из: {tech_literature_path}")
    success = loader.load_directory(tech_literature_path)
    
    if success:
        print("\n✅ Загрузка завершена успешно!")
        
        # Показываем статистику
        stats = loader.get_document_stats()
        if stats:
            print(f"\n📊 Статистика:")
            print(f"   Всего документов: {stats['total_documents']}")
            
            print(f"\n📂 Категории:")
            for category, count in stats['categories']:
                print(f"   {category}: {count} документов")
                
            print(f"\n📄 Типы файлов:")
            for file_type, count in stats['file_types']:
                print(f"   {file_type}: {count} файлов")
                
        print(f"\n🔍 Теперь вы можете искать документы через Rubin AI!")
        print(f"   База данных: {loader.db_path}")
        
    else:
        print("❌ Ошибка при загрузке документов")

if __name__ == "__main__":
    main()

